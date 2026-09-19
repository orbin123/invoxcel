import io
import json
import tempfile
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from tests.base import AuthenticatedTestCase

from apps.invoices.models import Invoice
from apps.processing.models import AnalysisBatch, BatchColumn
from apps.processing.tests import FIXTURE
from apps.templates.models import STANDARD_SUMMARY_COLUMNS, Template, TemplateColumn
from services.azure_document_intelligence import map_response


class OutputFlowTests(AuthenticatedTestCase):
    def setUp(self):
        super().setUp()
        media = tempfile.TemporaryDirectory()
        self.addCleanup(media.cleanup)
        settings = override_settings(MEDIA_ROOT=media.name)
        settings.enable()
        self.addCleanup(settings.disable)
        self.provider = patch("apps.processing.services.get_extraction_client")
        self.client_factory = self.provider.start()
        self.addCleanup(self.provider.stop)
        self.provider_client = self.client_factory.return_value
        self.provider_client.provider = "fixture"
        self.provider_client.model = "prebuilt-invoice"
        self.provider_client.extract.side_effect = lambda document: map_response(FIXTURE.read_text())

    def upload(self, template="", count=1):
        response = self.client.post(reverse("uploads:upload"), {
            "document": [SimpleUploadedFile(f"sample-{i}.pdf", b"%PDF-1.4\n%%EOF") for i in range(count)],
            "template": template,
        })
        self.assertEqual(response.status_code, 302)
        return AnalysisBatch.objects.latest("pk")

    def export(self, batch, kind="json"):
        Invoice.objects.filter(analysis__batch=batch).update(reviewed_at=timezone.now())
        return self.client.post(reverse("exports:download", args=[batch.pk, kind]))

    def test_standard_headers_exact_order_nulls_and_xlsx(self):
        template = Template.objects.get(is_standard=True)
        expected = [label for _, label, _ in STANDARD_SUMMARY_COLUMNS]
        self.assertEqual(list(template.columns.filter(dataset="summary").values_list("label", flat=True)), expected)
        batch = self.upload(template.pk)
        invoice = batch.document_analyses.get().invoice
        self.assertEqual(invoice.values["subtotal"], "100.10")
        self.assertEqual(invoice.values["taxable_amount"], "100.10")
        self.assertEqual(invoice.values["calculated_total"], "110.11")
        self.assertEqual(invoice.values["total_difference"], "0.00")
        self.assertIsNone(invoice.values["vendor_gstin"])
        self.assertEqual(invoice.values["verification_status"], "Review Required")
        workbook = load_workbook(io.BytesIO(self.export(batch, "xlsx").content))
        rows = list(workbook["Invoice Summary"].values)
        self.assertEqual(list(rows[0]), expected)
        self.assertEqual(len(rows[1]), 25)
        self.assertTrue(rows[1][0].startswith("INV"))
        self.assertEqual(rows[1][17], 110.11)
        self.assertIsNone(rows[1][4])

    def test_none_discovers_unknown_scalar_fields_and_item_columns(self):
        payload = json.loads(FIXTURE.read_text())
        fields = payload["analyzeResult"]["documents"][0]["fields"]
        fields["ProjectCode"] = {"type": "string", "valueString": "P-102"}
        fields["AmountDue"] = {"type": "currency", "valueCurrency": {"amount": 12.25}}
        self.provider_client.extract.side_effect = lambda document: map_response(json.dumps(payload))
        batch = self.upload()
        self.assertIsNone(batch.template_id)
        self.assertTrue(batch.columns.filter(key="project_code", dataset="summary").exists())
        self.assertEqual(batch.columns.get(key="amount_due").data_type, "currency")
        self.assertFalse(batch.columns.filter(key="vendor_gstin").exists())
        self.assertTrue(batch.columns.filter(key="quantity", dataset="line_item").exists())
        response = self.export(batch)
        self.assertIsNone(response.json()["template"]["name"])
        self.assertEqual(response.json()["invoices"][0]["summary"]["project_code"], "P-102")
        self.assertEqual(self.export(batch, "xlsx").status_code, 200)

    def test_custom_aliases_merge_and_unknown_columns_stay_null(self):
        template = Template.objects.create(name="Company import", user=self.user)
        for position, label in enumerate(("INVOICE NO", "Vendor-name", "Project code")):
            TemplateColumn.objects.create(template=template, dataset="summary", key=f"custom_{position}", label=label, data_type="text", position=position)
        batch = self.upload(template.pk)
        values = batch.document_analyses.get().invoice.values
        self.assertEqual(values["custom_0"], "DEMO-001")
        self.assertEqual(values["custom_1"], "Demo Office Supplies")
        self.assertIsNone(values["custom_2"])
        response = self.export(batch).json()
        self.assertIsNone(response["invoices"][0]["summary"]["custom_2"])

    def test_hide_empty_applies_across_batch_preserves_zero_and_all_export_formats(self):
        template = Template.objects.get(is_standard=True)
        batch = self.upload(template.pk, count=2)
        response = self.client.post(reverse("workspace:columns", args=[batch.pk]), {
            "columns": list(batch.columns.values_list("pk", flat=True)),
            "missing_columns": "hide",
        })
        self.assertEqual(response.status_code, 302)
        invoices = list(Invoice.objects.filter(analysis__batch=batch))
        invoices[1].values["vendor_gstin"] = "32ABCDE1234F1Z5"
        invoices[1].save()
        response = self.client.get(reverse("workspace:detail", args=[batch.pk]))
        visible = {c.key for c in response.context["summary_columns"]}
        self.assertIn("vendor_gstin", visible)
        self.assertIn("total_difference", visible)
        self.assertNotIn("customer_gstin", visible)
        data = self.export(batch).json()
        self.assertIn("vendor_gstin", data["invoices"][0]["summary"])
        self.assertIsNone(data["invoices"][0]["summary"]["vendor_gstin"])
        self.assertNotIn("customer_gstin", data["invoices"][0]["summary"])
        workbook = load_workbook(io.BytesIO(self.export(batch, "xlsx").content))
        self.assertNotIn("Customer GSTIN", list(workbook["Invoice Summary"].values)[0])
        self.assertNotIn("Customer GSTIN", self.export(batch, "summary.csv").content.decode())
        self.assertEqual(self.provider_client.extract.call_count, 2)

    def test_missing_column_choice_is_only_in_template_workspace(self):
        home = self.client.get(reverse("uploads:upload"))
        self.assertNotContains(home, 'name="missing_columns"')
        self.assertNotIn("missing_columns", home.context["upload_form"].fields)
        batch = self.upload(Template.objects.get(is_standard=True).pk)
        self.assertFalse(batch.hide_empty_columns)
        workspace = self.client.get(reverse("workspace:detail", args=[batch.pk]))
        self.assertContains(workspace, "Missing template columns")
        self.assertContains(workspace, 'name="missing_columns"')
        auto_batch = self.upload()
        workspace = self.client.get(reverse("workspace:detail", args=[auto_batch.pk]))
        self.assertNotContains(workspace, 'name="missing_columns"')
        self.assertNotContains(workspace, "Missing template columns")
        self.assertContains(workspace, "Delete column")
        self.assertContains(workspace, "Apply visibility")
        response = self.client.post(reverse("workspace:columns", args=[auto_batch.pk]), {
            "columns": list(auto_batch.columns.values_list("pk", flat=True)),
            "missing_columns": "hide",
        })
        self.assertEqual(response.status_code, 400)
        auto_batch.refresh_from_db()
        self.assertFalse(auto_batch.hide_empty_columns)

    def test_old_upload_option_cannot_enable_hiding_and_none_ignores_legacy_setting(self):
        response = self.client.post(reverse("uploads:upload"), {
            "document": SimpleUploadedFile("sample.pdf", b"%PDF-1.4\n%%EOF"),
            "template": "", "missing_columns": "hide",
        })
        self.assertEqual(response.status_code, 302)
        batch = AnalysisBatch.objects.latest("pk")
        self.assertFalse(batch.hide_empty_columns)
        batch.hide_empty_columns = True
        batch.save()
        column = batch.columns.get(key="vendor_name")
        invoice = batch.document_analyses.get().invoice
        invoice.values["vendor_name"] = None
        invoice.save()
        workspace = self.client.get(reverse("workspace:detail", args=[batch.pk]))
        self.assertIn(column, workspace.context["summary_columns"])
        self.assertIn("vendor_name", self.export(batch).json()["invoices"][0]["summary"])

    def test_rename_delete_columns_keep_template_and_export_headers_literal(self):
        template = Template.objects.get(is_standard=True)
        batch = self.upload(template.pk)
        column = batch.columns.get(key="vendor_name")
        url = reverse("workspace:edit_column", args=[batch.pk, column.pk])
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertEqual(self.client.post(url, {"action": "rename", "label": "=1+1"}).status_code, 302)
        workbook = load_workbook(io.BytesIO(self.export(batch, "xlsx").content))
        cell = next(c for c in workbook["Invoice Summary"][1] if c.value == "=1+1")
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(template.columns.get(key="vendor_name").label, "Vendor Name")
        self.assertEqual(self.client.post(url, {"action": "delete"}).status_code, 302)
        self.assertFalse(batch.columns.filter(pk=column.pk).exists())
        self.assertNotIn("vendor_name", self.export(batch).json()["invoices"][0]["summary"])
        self.assertEqual(self.provider_client.extract.call_count, 1)

    def test_delete_invoice_row_removes_items_from_exports_and_retains_source(self):
        batch = self.upload(count=2)
        analysis = batch.document_analyses.first()
        document = analysis.document
        invoice_id = analysis.invoice.pk
        url = reverse("workspace:delete_row", args=[batch.pk, analysis.pk])
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertEqual(self.client.post(url).status_code, 302)
        self.assertFalse(Invoice.objects.filter(pk=invoice_id).exists())
        self.assertTrue(document.file.storage.exists(document.file.name))
        self.assertEqual(len(self.export(batch).json()["invoices"]), 1)
        other_batch = self.upload()
        column = other_batch.columns.first()
        self.assertEqual(self.client.post(reverse("workspace:edit_column", args=[batch.pk, column.pk]), {"action": "delete"}).status_code, 404)
        self.assertEqual(self.client.post(reverse("workspace:delete_row", args=[batch.pk, other_batch.document_analyses.get().pk])).status_code, 404)

    def test_company_sample_totals_discount_once_and_missing_tax_not_invented(self):
        from apps.invoices.validation import update_derived_values
        batch = self.upload(Template.objects.get(is_standard=True).pk)
        invoice = batch.document_analyses.get().invoice
        invoice.values.update(subtotal="25000", discount="500", cgst="2205", sgst="2205", igst="0", other_charges="250", grand_total="29160")
        update_derived_values(invoice)
        self.assertEqual(invoice.values["taxable_amount"], "24500")
        self.assertEqual(invoice.values["calculated_total"], "29160")
        self.assertEqual(invoice.values["total_difference"], "0")
        for key in ("cgst", "sgst", "igst", "tax_amount"):
            invoice.values[key] = None
        update_derived_values(invoice)
        self.assertIsNone(invoice.values["calculated_total"])

    def test_custom_numeric_alias_remains_typed_after_rename_and_edits_update_totals(self):
        template = Template.objects.create(name="Amounts", user=self.user)
        for position, label in enumerate(("Subtotal", "Discount", "Grand total")):
            TemplateColumn.objects.create(template=template, dataset="summary", key=f"custom_{position}", label=label, data_type="text", position=position)
        batch = self.upload(template.pk)
        invoice = batch.document_analyses.get().invoice
        column = batch.columns.get(key="custom_0")
        self.assertEqual(column.data_type, "currency")
        self.client.post(reverse("workspace:edit_column", args=[batch.pk, column.pk]), {"action": "rename", "label": "Base amount"})
        item = invoice.line_items.get()
        data = {"revision": invoice.updated_at.isoformat(), "action": "review", "summary-custom_0": "120", "summary-custom_1": "10", "summary-custom_2": "120.01", "items-TOTAL_FORMS": "1", "items-INITIAL_FORMS": "1", "items-0-row_id": item.pk}
        response = self.client.post(reverse("workspace:detail", args=[batch.pk]), data)
        self.assertEqual(response.status_code, 302)
        invoice.refresh_from_db()
        self.assertEqual(invoice.values["subtotal"], "120")
        self.assertEqual(invoice.values["taxable_amount"], "110")
        self.assertEqual(invoice.values["calculated_total"], "120.01")
        self.assertEqual(invoice.values["verification_status"], "Approved")
        data.update({"revision": invoice.updated_at.isoformat(), "summary-custom_0": "unsafe"})
        response = self.client.post(reverse("workspace:detail", args=[batch.pk]), data)
        self.assertContains(response, "Enter a finite number")
        invoice.refresh_from_db()
        self.assertEqual(invoice.values["subtotal"], "120")

    def test_later_invoice_can_populate_previously_missing_custom_column(self):
        template = Template.objects.create(name="Projects", user=self.user)
        TemplateColumn.objects.create(template=template, dataset="summary", key="custom_project_code", label="Project code", data_type="text", position=0)
        payload = json.loads(FIXTURE.read_text())
        payload["analyzeResult"]["documents"][0]["fields"]["ProjectCode"] = {"type": "string", "valueString": "P-999"}
        self.provider_client.extract.side_effect = [map_response(FIXTURE.read_text()), map_response(json.dumps(payload))]
        batch = self.upload(template.pk, count=2)
        values = list(Invoice.objects.filter(analysis__batch=batch).order_by("pk").values_list("values", flat=True))
        self.assertIsNone(values[0]["custom_project_code"])
        self.assertEqual(values[1]["custom_project_code"], "P-999")

    def test_tiff_upload_has_local_multi_page_preview(self):
        from PIL import Image
        output = io.BytesIO()
        first, second = Image.new("RGB", (100, 100), "white"), Image.new("RGB", (100, 100), "gray")
        first.save(output, format="TIFF", save_all=True, append_images=[second])
        result = map_response(FIXTURE.read_text())
        result.page_count = 2
        self.provider_client.extract.side_effect = lambda document: result
        response = self.client.post(reverse("uploads:upload"), {"document": SimpleUploadedFile("scan.tiff", output.getvalue())})
        self.assertEqual(response.status_code, 302)
        document = AnalysisBatch.objects.latest("pk").document_analyses.get().document
        self.assertEqual(document.content_type, "image/tiff")
        preview = self.client.get(reverse("uploads:preview", args=[document.pk, 2]))
        self.assertEqual(preview.status_code, 200)
        self.assertTrue(preview.content.startswith(b"\x89PNG"))
        self.assertEqual(self.client.get(reverse("uploads:preview", args=[document.pk, 3])).status_code, 404)
