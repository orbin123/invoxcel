import csv
import io
import tempfile
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.urls import reverse
from openpyxl import load_workbook
from pypdf import PdfWriter

from tests.base import AuthenticatedTestCase

from apps.invoices.models import Invoice
from apps.processing.models import AnalysisBatch, DocumentAnalysis
from apps.processing.services import extract_batch
from apps.processing.tests import FIXTURE
from apps.uploads.models import UploadedDocument
from services.azure_document_intelligence import ExtractionError, map_response


class FixtureClient:
    provider = "fixture"
    model = "prebuilt-invoice"

    def extract(self, document):
        return map_response(FIXTURE.read_text())


class WorkflowTests(AuthenticatedTestCase):
    def setUp(self):
        super().setUp()
        self.media = tempfile.TemporaryDirectory()
        self.addCleanup(self.media.cleanup)
        self.settings_override = override_settings(MEDIA_ROOT=self.media.name)
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.provider = patch("apps.processing.services.get_extraction_client", return_value=FixtureClient())
        self.mock_provider = self.provider.start()
        self.addCleanup(self.provider.stop)
        response = self.client.post(reverse("uploads:upload"), {"document": SimpleUploadedFile("test.pdf", b"%PDF-1.4\n%%EOF")})
        self.batch = AnalysisBatch.objects.latest("pk")
        self.url = reverse("workspace:detail", args=[self.batch.pk])
        self.assertRedirects(response, self.url)
        self.invoice = Invoice.objects.get(analysis__batch=self.batch)

    def edit_data(self, action="review"):
        self.invoice.refresh_from_db()
        items = list(self.invoice.line_items.all())
        data = {"action": action, "revision": self.invoice.updated_at.isoformat(), "items-TOTAL_FORMS": str(len(items)), "items-INITIAL_FORMS": str(len(items)), "items-MIN_NUM_FORMS": "0", "items-MAX_NUM_FORMS": "200"}
        data.update({f"summary-{c.key}": self.invoice.values.get(c.key) or "" for c in self.batch.columns.filter(dataset="summary")})
        for index, item in enumerate(items):
            data[f"items-{index}-row_id"] = item.pk
            data.update({f"items-{index}-{c.key}": item.values.get(c.key) or "" for c in self.batch.columns.filter(dataset="line_item")})
        return data

    def test_upload_edit_review_and_all_exports_use_saved_data(self):
        analysis = self.invoice.analysis
        self.assertEqual(analysis.raw_response, FIXTURE.read_text())
        self.assertContains(self.client.get(self.url), "DEMO-001")
        data = self.edit_data()
        data["summary-vendor_name"] = "Corrected Supplier"
        data["items-0-description"] = "Corrected line"
        self.assertRedirects(self.client.post(self.url, data), self.url + f"?document={analysis.pk}")
        self.invoice.refresh_from_db()
        self.assertIsNotNone(self.invoice.reviewed_at)
        for format in ("summary.csv", "line-items.csv", "json", "xlsx"):
            response = self.client.post(reverse("exports:download", args=[self.batch.pk, format]))
            self.assertEqual(response.status_code, 200)
            if format == "xlsx":
                workbook = load_workbook(io.BytesIO(response.content))
                self.assertEqual(workbook.sheetnames, ["Invoice Summary", "Line Items", "Exceptions"])
                summary = list(workbook["Invoice Summary"].values)
                self.assertIn("Corrected Supplier", summary[1])
                self.assertIn(110.11, summary[1])
                self.assertIn("Corrected line", list(workbook["Line Items"].values)[1])
            elif format == "json":
                self.assertEqual(response.json()["invoices"][0]["summary"]["vendor_name"], "Corrected Supplier")
            else:
                rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
                self.assertIn("Corrected Supplier" if format == "summary.csv" else "Corrected line", rows[1])
        self.assertEqual(self.mock_provider.call_count, 1)

    def test_unreviewed_export_blocked_and_get_does_not_mutate(self):
        url = reverse("exports:download", args=[self.batch.pk, "json"])
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertRedirects(self.client.post(url), self.url)
        self.assertIsNone(Invoice.objects.get(pk=self.invoice.pk).reviewed_at)

    def test_output_preview_keeps_saved_line_items_separate_and_in_column_order(self):
        response = self.client.get(self.url)
        columns = response.context["item_columns"]
        items = list(self.invoice.line_items.all())
        rows = response.context["item_output_rows"]
        self.assertEqual(len(rows), len(items))
        for row, item in zip(rows, items):
            self.assertEqual(row["document_id"], self.invoice.analysis.document_id)
            self.assertEqual([cell["value"] for cell in row["cells"]],
                             [item.values.get(column.key) for column in columns])
        self.assertContains(response, 'id="export-format"')
        self.assertContains(response, "The Exceptions sheet includes source filenames")

    def test_excel_metadata_note_matches_existing_export_sheets(self):
        self.client.post(self.url, self.edit_data())
        response = self.client.post(reverse("exports:download", args=[self.batch.pk, "xlsx"]))
        workbook = load_workbook(io.BytesIO(response.content))
        for name in ("Invoice Summary", "Line Items"):
            headers = [cell.value for cell in workbook[name][1]]
            self.assertNotIn("Source file", headers)
            self.assertNotIn("Review status", headers)
        self.assertIn("Source filename", [cell.value for cell in workbook["Exceptions"][1]])

    def test_invalid_edit_does_not_persist_any_change(self):
        data = self.edit_data()
        data.update({"summary-grand_total": "NaN", "summary-vendor_name": "Should not persist"})
        response = self.client.post(self.url, data)
        self.assertContains(response, "Enter a finite number")
        self.invoice.refresh_from_db()
        self.assertEqual(self.invoice.values["vendor_name"], "Demo Office Supplies")
        self.assertIsNone(self.invoice.reviewed_at)

    def test_save_without_review_resets_review_and_stale_revision_is_rejected(self):
        data = self.edit_data()
        self.client.post(self.url, data)
        response = self.client.post(self.url, data)
        self.assertContains(response, "changed in another tab")
        data = self.edit_data("save")
        self.client.post(self.url, data)
        self.invoice.refresh_from_db()
        self.assertIsNone(self.invoice.reviewed_at)

    def test_add_and_delete_items_and_reject_foreign_row_ids(self):
        data = self.edit_data()
        data["items-0-row_id"] = 999999
        self.assertEqual(self.client.post(self.url, data).status_code, 400)
        data = self.edit_data()
        data.update({"items-TOTAL_FORMS": "2", "items-0-DELETE": "on", "items-1-description": "Added row", "items-1-quantity": "3", "items-1-unit_price": "10", "items-1-amount": "30"})
        self.assertEqual(self.client.post(self.url, data).status_code, 302)
        self.assertEqual(self.invoice.line_items.count(), 1)
        self.assertEqual(self.invoice.line_items.get().values["description"], "Added row")

    def test_column_visibility_does_not_reextract_or_erase_canonical_data(self):
        selected = self.batch.columns.filter(key="invoice_number").get()
        self.client.post(reverse("workspace:columns", args=[self.batch.pk]), {"columns": [selected.pk]})
        self.assertEqual(self.mock_provider.call_count, 1)
        self.invoice.refresh_from_db()
        self.assertEqual(self.invoice.values["grand_total"], "110.11")
        self.assertEqual(self.batch.columns.filter(is_selected=True).count(), 1)

    def test_formula_text_is_literal_in_xlsx_and_escaped_in_csv(self):
        data = self.edit_data()
        data["summary-vendor_name"] = '=HYPERLINK("https://example.com")'
        self.client.post(self.url, data)
        csv_response = self.client.post(reverse("exports:download", args=[self.batch.pk, "summary.csv"]))
        rows = list(csv.reader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
        self.assertIn("'" + data["summary-vendor_name"], rows[1])
        response = self.client.post(reverse("exports:download", args=[self.batch.pk, "xlsx"]))
        workbook = load_workbook(io.BytesIO(response.content))
        vendor_cell = next(c for c in workbook["Invoice Summary"][2] if c.value == data["summary-vendor_name"])
        self.assertEqual(vendor_cell.data_type, "s")

    def test_failed_upload_retains_source_can_retry_and_does_not_repeat_completed_work(self):
        with patch.object(FixtureClient, "extract", side_effect=ExtractionError("Azure is unavailable.")):
            response = self.client.post(reverse("uploads:upload"), {"document": SimpleUploadedFile("retry.pdf", b"%PDF-1.4\n%%EOF")})
        batch = AnalysisBatch.objects.latest("pk")
        analysis = batch.document_analyses.get()
        self.assertEqual(analysis.status, "failed")
        self.assertTrue(analysis.document.file.storage.exists(analysis.document.file.name))
        self.assertContains(self.client.get(response.url), "Retry extraction")
        retry_url = reverse("workspace:retry", args=[batch.pk, analysis.pk])
        self.assertEqual(self.client.get(retry_url).status_code, 405)
        self.assertEqual(self.client.post(retry_url).status_code, 302)
        analysis.refresh_from_db()
        self.assertEqual(analysis.status, "complete")
        with patch.object(FixtureClient, "extract") as provider:
            extract_batch(batch, FixtureClient())
            provider.assert_not_called()
        self.assertEqual(batch.document_analyses.get().invoice.line_items.count(), 1)

    def test_pdf_source_preview_renders_and_invalid_page_returns_404(self):
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        pdf = io.BytesIO()
        writer.write(pdf)
        self.invoice.analysis.document.file.save("preview.pdf", SimpleUploadedFile("preview.pdf", pdf.getvalue()))
        source_url = reverse("uploads:source", args=[self.invoice.analysis.document_id])
        self.assertEqual(self.client.get(source_url)["X-Frame-Options"], "SAMEORIGIN")
        response = self.client.get(reverse("uploads:preview", args=[self.invoice.analysis.document_id, 1]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"\x89PNG"))
        self.assertEqual(self.client.get(reverse("uploads:preview", args=[self.invoice.analysis.document_id, 3])).status_code, 404)

    def test_document_navigator_switches_source_and_extracted_output(self):
        analyses = [self.invoice.analysis]
        for filename, invoice_number in (("second.pdf", "SECOND-002"), ("third.jpeg", "THIRD-003")):
            document = UploadedDocument.objects.create(
                user=self.user,
                original_filename=filename,
                file=SimpleUploadedFile(filename, b"invoice image"),
                content_type="image/jpeg" if filename.endswith("jpeg") else "application/pdf",
                size_bytes=13,
            )
            analysis = DocumentAnalysis.objects.create(batch=self.batch, document=document, status="complete")
            Invoice.objects.create(analysis=analysis, values={"invoice_number": invoice_number})
            analyses.append(analysis)

        first = self.client.get(self.url)
        self.assertContains(first, "Invoice 1 of 3")
        self.assertContains(first, f'?document={analyses[1].pk}')
        self.assertContains(first, '<span class="previous-document" aria-disabled="true">← Previous</span>', html=True)

        second = self.client.get(f"{self.url}?document={analyses[1].pk}")
        self.assertEqual(second.context["selected"], analyses[1])
        self.assertContains(second, "second.pdf")
        self.assertContains(second, "SECOND-002")
        self.assertContains(second, "Invoice 2 of 3")
        self.assertContains(second, f'?document={analyses[0].pk}')
        self.assertContains(second, f'?document={analyses[2].pk}')

        third = self.client.get(f"{self.url}?document={analyses[2].pk}")
        self.assertContains(third, "THIRD-003")
        self.assertContains(third, '<span class="next-document" aria-disabled="true">Next →</span>', html=True)

    def test_failed_documents_are_in_export_exceptions_and_unreviewed_success_blocks_batch(self):
        document = UploadedDocument.objects.create(user=self.user, original_filename="failed.pdf", file=SimpleUploadedFile("failed.pdf", b"%PDF-1.4\n%%EOF"), content_type="application/pdf", size_bytes=14)
        DocumentAnalysis.objects.create(batch=self.batch, document=document, status="failed", error_message="Azure unavailable")
        self.client.post(self.url, self.edit_data())
        response = self.client.post(reverse("exports:download", args=[self.batch.pk, "json"]))
        self.assertEqual(len(response.json()["invoices"]), 1)
        self.assertEqual(response.json()["exceptions"][0][1], "failed.pdf")
        self.assertEqual(response.json()["exceptions"][0][3], "Azure unavailable")
        analysis = DocumentAnalysis.objects.create(batch=self.batch, document=UploadedDocument.objects.create(user=self.user, original_filename="second.pdf", file=SimpleUploadedFile("second.pdf", b"%PDF-1.4\n%%EOF"), content_type="application/pdf", size_bytes=14), status="complete")
        Invoice.objects.create(analysis=analysis)
        self.assertRedirects(self.client.post(reverse("exports:download", args=[self.batch.pk, "json"])), self.url)

    def test_raw_malformed_output_is_saved_for_failed_extraction(self):
        with patch.object(FixtureClient, "extract", side_effect=ExtractionError("Unusable invoice", raw_response='{"analyzeResult": {}}')):
            self.client.post(reverse("uploads:upload"), {"document": SimpleUploadedFile("bad.pdf", b"%PDF-1.4\n%%EOF")})
        analysis = DocumentAnalysis.objects.latest("pk")
        self.assertEqual(analysis.status, "failed")
        self.assertEqual(analysis.raw_response, '{"analyzeResult": {}}')
        self.assertFalse(Invoice.objects.filter(analysis=analysis).exists())

    def test_programmatic_checks_include_hidden_optional_template_values_and_refresh(self):
        from apps.templates.models import Template
        from apps.processing.models import BatchColumn
        self.batch.template = Template.objects.filter(is_standard=True).first()
        self.batch.hide_empty_columns = True
        self.batch.save()
        column = BatchColumn.objects.create(batch=self.batch, dataset="summary", key="po_number", label="PO No.",
                                          data_type="text", position=90, is_selected=False)
        self.invoice.provenance["vendor_name"] = {"confidence": 0.5}
        self.invoice.values["po_number"] = None
        self.invoice.save()
        response = self.client.get(self.url)
        self.assertContains(response, "Azure extraction checks")
        self.assertContains(response, "Programmatic checks")
        self.assertTrue(any(f["field"] == "vendor_name" for f in response.context["azure_findings"]))
        self.assertTrue(any(f["field"] == column.key for f in response.context["programmatic_findings"]))
        self.client.post(self.url, self.edit_data())
        response = self.client.post(reverse("exports:download", args=[self.batch.pk, "json"]))
        self.assertTrue(any(f["field"] == column.key for f in response.json()["invoices"][0]["validation"]))
        self.invoice.refresh_from_db()
        self.invoice.values["po_number"] = "PO-123"
        self.invoice.save()
        response = self.client.get(self.url)
        self.assertFalse(any(f["field"] == column.key for f in response.context["programmatic_findings"]))
