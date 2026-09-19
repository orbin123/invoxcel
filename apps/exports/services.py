import csv
import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill


def csv_text(value):
    text = "" if value is None else str(value)
    # Spreadsheet applications interpret these prefixes even inside quoted CSV cells.
    if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")):
        return "'" + text
    return text


def render_csv(headers, rows):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow([csv_text(value) for value in headers])
    writer.writerows([[csv_text(value) for value in row] for row in rows])
    return "\ufeff" + output.getvalue()


def includes_invoice_id(columns):
    return any((c.value_key or c.key) == "invoice_id" for c in columns)


def export_headers(columns, summary=False):
    return ([] if summary and includes_invoice_id(columns) else ["Document ID"]) + [c.label for c in columns]


def datasets(invoices, summary_columns, item_columns):
    summary = [([] if includes_invoice_id(summary_columns) else [invoice.analysis.document_id]) + [invoice.values.get(c.key) for c in summary_columns] for invoice in invoices]
    items = [[invoice.analysis.document_id] + [item.values.get(c.key) for c in item_columns] for invoice in invoices for item in invoice.line_items.all()]
    return summary, items


def render_xlsx(summary_columns, item_columns, summary, items, exceptions):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, columns, rows in (("Invoice Summary", summary_columns, summary), ("Line Items", item_columns, items)):
        sheet = workbook.create_sheet(title)
        sheet.append([ILLEGAL_CHARACTERS_RE.sub("", label) for label in export_headers(columns, summary=title == "Invoice Summary")])
        for cell in sheet[1]:
            cell.data_type = "s"
        for row in rows:
            has_id = title != "Invoice Summary" or not includes_invoice_id(columns)
            sheet.append([row[0]] if has_id else [None])
            for index, (column, value) in enumerate(zip(columns, row[1:] if has_id else row), 2 if has_id else 1):
                cell = sheet.cell(sheet.max_row, index)
                if value is None:
                    continue
                if column.data_type in ("number", "currency"):
                    cell.value = Decimal(value)
                    cell.number_format = '#,##0.00########' if column.data_type == "currency" else '0.########'
                elif column.data_type == "date":
                    cell.value = date.fromisoformat(value)
                    cell.number_format = "yyyy-mm-dd"
                elif column.data_type == "boolean":
                    cell.value = value
                else:
                    cell.value = ILLEGAL_CHARACTERS_RE.sub("", str(value))
                    cell.data_type = "s"  # Never turn document text into executable formulas.
    sheet = workbook.create_sheet("Exceptions")
    sheet.append(["Document ID", "Source filename", "Field", "Finding"])
    for row in exceptions:
        sheet.append([row[0]])
        for index, value in enumerate(row[1:], 2):
            cell = sheet.cell(sheet.max_row, index, ILLEGAL_CHARACTERS_RE.sub("", str(value)))
            cell.data_type = "s"
    for sheet in workbook:
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8EFEA")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(50, max(16, max(len(str(cell.value or "")) for cell in column) + 2))
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
